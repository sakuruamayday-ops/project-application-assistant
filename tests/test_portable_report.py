import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


class PortableReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.repository = Path(__file__).resolve().parents[1]
        cls.skills = cls.repository / "skills"
        cls.tax = cls.skills / "manufacturing-tax-risk-analysis"
        cls.runtime = cls.skills / "_runtime" / "gongchuang-branding"

    def test_runtime_guidance_does_not_make_the_full_example_a_prerequisite(self):
        skill = (self.tax / "SKILL.md").read_text(encoding="utf-8")
        schema = (self.tax / "references/report-input-schema.md").read_text(encoding="utf-8")
        self.assertIn("首次生成不得读取或复制该示例", skill)
        self.assertIn("不要在首次运行前读取或复制同目录的示例文件", schema)
        self.assertIn("不得重新读取完整 HTML", skill)
        self.assertIn("KPI 卡片不超过4项", schema)
        self.assertIn("不得自行计算三年复合增长率", schema)
        self.assertIn("不得把各年净利润相加后写成累计净利润", schema)
        self.assertIn("不要先用不确定的相对路径试错", skill)
        self.assertIn("不得只写“完成”“已完成”“通过”“达标”“待完成”", schema)

    def generate_metrics(self, directory: Path) -> Path:
        facts = directory / "facts.json"
        metrics = directory / "metrics.json"
        subprocess.run(
            [
                sys.executable,
                str(self.tax / "scripts/calculate_metrics.py"),
                str(self.tax / "references/metrics-input.example.json"),
                str(facts),
                "--metrics-output",
                str(metrics),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        return metrics

    def test_prepared_input_renders_calculator_fields_and_preserves_existing_edits(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            metrics = self.generate_metrics(root)
            draft = root / "draft.json"
            command = [sys.executable, str(self.tax / "scripts/generate_report_html.py"), str(draft), "--metrics-json", str(metrics)]
            prepared = subprocess.run([*command, "--prepare-input"], capture_output=True, text=True)
            self.assertEqual(prepared.returncode, 0, prepared.stderr)
            original = draft.read_bytes()
            data = json.loads(original)
            self.assertIn("{{financial.", data["financial_overview"]["rows"][0]["values"][0])
            self.assertEqual(data["policies"], [])
            self.assertEqual(data["calculations"], [])
            self.assertEqual(data["risk_level"], "待判断")
            repeat = subprocess.run([*command, "--prepare-input"], capture_output=True, text=True)
            self.assertNotEqual(repeat.returncode, 0)
            self.assertEqual(draft.read_bytes(), original)
            output = root / "report.html"
            result = subprocess.run([sys.executable, str(self.tax / "scripts/generate_report_html.py"), str(draft), str(output), "--metrics-json", str(metrics)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            rendered = output.read_text(encoding="utf-8")
            self.assertNotIn("enterprise-financial-facts/v1", rendered)
            self.assertIn("判断边界：", rendered)
            self.assertNotIn("证据闸门：", rendered)
            self.assertTrue(all("年" in row["source"] for row in json.loads(metrics.read_text())["report_rows"]))
            self.assertEqual(rendered.count('<section class="page'), 17)
            self.assertNotIn("{{financial.", rendered)
            values = json.loads(metrics.read_text(encoding="utf-8"))["display_values"]
            latest = sorted(values)[-1]
            self.assertIn(values[latest]["facts"]["other_receivables"], rendered)
            self.assertIn("经营现金流与净利润之比", rendered)
            self.assertNotIn("余额集中于少数对象", rendered)

    def test_generator_outputs_exactly_seventeen_page_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            output = directory_path / "report.html"
            metrics = self.generate_metrics(directory_path)
            result = subprocess.run(
                [
                    sys.executable,
                    str(self.tax / "scripts/generate_report_html.py"),
                    str(self.tax / "references/report-data.example.json"),
                    str(output),
                    "--metrics-json",
                    str(metrics),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            payload = json.loads(result.stdout)
            self.assertEqual(payload["pages"], 17)
            self.assertEqual(
                output.read_text(encoding="utf-8").count('<section class="page'),
                17,
            )
            html = output.read_text(encoding="utf-8")
            self.assertEqual(html.count('class="gongchuang-page-watermark"'), 17)
            self.assertEqual(html.count('class="gongchuang-document-header"'), 17)
            self.assertEqual(html.count('class="gongchuang-cover-signature"'), 1)
            self.assertIn('<meta name="author"', html)
            self.assertIn('.cover h2{color:#fff8e8}', html)
            self.assertIn('.roadmap-table th:first-child,.roadmap-table td:first-child{white-space:nowrap;width:24mm}', html)

    def test_overview_preserves_existing_year_suffixes(self):
        spec = importlib.util.spec_from_file_location("tax_report", self.tax / "scripts/generate_report_html.py")
        generator = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(generator)
        data = json.loads((self.tax / "references/report-data.example.json").read_text(encoding="utf-8"))
        data["financial_overview"]["years"] = ["2023", "2024年", "2025年度"]
        html = generator.render_overview(data)
        for label in ("2023年：", "2024年：", "2025年度："):
            self.assertIn(label, html)
        self.assertNotIn("年年", html)
        self.assertNotIn("年度年", html)

    def test_financial_references_resolve_by_field_without_modifying_editor_input(self):
        spec = importlib.util.spec_from_file_location("tax_report", self.tax / "scripts/generate_report_html.py")
        generator = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(generator)
        original = {"risks": [{"fact": "其他应收款{{financial.2025.facts.other_receivables}}，占资产{{financial.2025.metrics.other_ar_to_assets}}。"}]}
        metrics = {"display_values": {"2025": {
            "facts": {"other_receivables": "25.00万元"},
            "metrics": {"other_ar_to_assets": "1.47%"},
        }}}
        resolved = generator.resolve_financial_references(original, metrics)
        self.assertEqual(resolved["risks"][0]["fact"], "其他应收款25.00万元，占资产1.47%。")
        self.assertIn("{{financial.", original["risks"][0]["fact"])
        with self.assertRaisesRegex(ValueError, "unknown financial reference"):
            generator.resolve_financial_references("{{financial.2024.facts.other_receivables}}", metrics)

    def test_generator_cli_fills_financial_references_before_validation_and_render(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            metrics_path = self.generate_metrics(directory_path)
            metrics = json.loads(metrics_path.read_text(encoding="utf-8"))
            year = sorted(metrics["display_values"])[-1]
            data = json.loads((self.tax / "references/report-data.example.json").read_text(encoding="utf-8"))
            data["risks"][0]["fact"] = "其他应收款{{financial." + year + ".facts.other_receivables}}。"
            source = directory_path / "report-data.json"
            source.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            output = directory_path / "report.html"
            subprocess.run([
                sys.executable, str(self.tax / "scripts/generate_report_html.py"),
                str(source), str(output), "--metrics-json", str(metrics_path),
            ], check=True, capture_output=True, text=True)
            html = output.read_text(encoding="utf-8")
            self.assertIn("其他应收款" + metrics["display_values"][year]["facts"]["other_receivables"] + "。", html)
            self.assertNotIn("{{financial.", html)
            self.assertIn("{{financial.", source.read_text(encoding="utf-8"))

    def test_generator_and_delivery_contract_use_the_same_visible_structure(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            output = directory_path / "report.html"
            metrics = self.generate_metrics(directory_path)
            subprocess.run(
                [
                    sys.executable,
                    str(self.tax / "scripts/generate_report_html.py"),
                    str(self.tax / "references/report-data.example.json"),
                    str(output),
                    "--metrics-json",
                    str(metrics),
                ],
                check=True,
            )
            html = output.read_text(encoding="utf-8")
            contracts = json.loads((self.skills / "delivery-contracts.json").read_text(encoding="utf-8"))
            profile = contracts["delivery_profiles"]["manufacturing-tax-risk-report"]
            cursor = 0
            for section in profile["required_sections"]:
                cursor = html.index(section, cursor) + len(section)
            for table_rule in profile["required_tables"]:
                self.assertIn(table_rule["id"], html)
                for column in table_rule["required_columns"]:
                    self.assertIn(column, html)
            for expected in (
                "2024年营业收入同比增长率",
                "2025年研发费用率",
            ):
                self.assertIn(expected, html)

    def test_source_rows_are_not_rejected_or_silently_dropped(self):
        spec = importlib.util.spec_from_file_location("tax_report", self.tax / "scripts/generate_report_html.py")
        generator = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(generator)
        data = json.loads((self.tax / "references/report-data.example.json").read_text(encoding="utf-8"))
        rows = [{"indicator": f"指标{index:02d}", "formula": "A/B", "result": "待核验", "source": f"来源{index:02d}"} for index in range(19)]
        metrics = {"schema": "manufacturing-tax-risk-metrics/v1", "company": data["company"], "report_rows": rows[:9]}
        data["calculations"] = rows[9:]
        self.assertEqual(generator.validate_metrics(metrics, data["company"]), rows[:9])
        html = generator.render(data, metrics, "")
        self.assertEqual(html.count('<section class="page'), 18)
        self.assertEqual(html.count('class="gongchuang-page-watermark"'), 18)
        for row in rows:
            self.assertEqual(html.count(f'<td>{row["indicator"]}</td>'), 1)
            self.assertIn(f'<td>{row["source"]}</td>', html)

    def test_generator_rejects_structured_values_in_all_human_text_lists(self):
        mutations = (
            (
                "section action",
                lambda data: data["sections"]["profitability"]["actions"].__setitem__(0, {"text": "错误对象"}),
                "field must be text: root.sections.profitability.actions[0]",
            ),
            (
                "roadmap action",
                lambda data: data["roadmap"][0]["actions"].__setitem__(0, {"text": "错误对象"}),
                "field must be text: root.roadmap[0].actions[0]",
            ),
            (
                "p0 document",
                lambda data: data["p0_documents"].__setitem__(0, {"name": "错误对象"}),
                "field must be text: root.p0_documents[0]",
            ),
            (
                "missing document",
                lambda data: data["missing_documents"].__setitem__(0, {"name": "错误对象"}),
                "field must be text: root.missing_documents[0]",
            ),
            (
                "limitation",
                lambda data: data["limitations"].__setitem__(0, {"text": "错误对象"}),
                "field must be text: root.limitations[0]",
            ),
        )
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            metrics = self.generate_metrics(directory_path)
            example = json.loads(
                (self.tax / "references/report-data.example.json").read_text(encoding="utf-8")
            )
            for index, (label, mutate, expected) in enumerate(mutations):
                with self.subTest(label=label):
                    candidate = json.loads(json.dumps(example, ensure_ascii=False))
                    mutate(candidate)
                    input_path = directory_path / f"invalid-{index}.json"
                    input_path.write_text(json.dumps(candidate, ensure_ascii=False), encoding="utf-8")
                    result = subprocess.run(
                        [
                            sys.executable,
                            str(self.tax / "scripts/generate_report_html.py"),
                            str(input_path),
                            "--validate-only",
                            "--metrics-json",
                            str(metrics),
                        ],
                        capture_output=True,
                        text=True,
                    )
                    self.assertNotEqual(result.returncode, 0)
                    self.assertIn(expected, result.stderr)

    def test_generator_rejects_non_object_report_rows_before_rendering(self):
        mutations = (
            ("sources", "field must be an object: root.sources[0]"),
            ("executive_findings", "field must be an object: root.executive_findings[0]"),
            ("risks", "field must be an object: root.risks[0]"),
            ("calculations", "field must be an object: root.calculations[0]"),
            ("policies", "field must be an object: root.policies[0]"),
            ("monthly_indicators", "field must be an object: root.monthly_indicators[0]"),
        )
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            metrics = self.generate_metrics(directory_path)
            example = json.loads(
                (self.tax / "references/report-data.example.json").read_text(encoding="utf-8")
            )
            for index, (field, expected) in enumerate(mutations):
                with self.subTest(field=field):
                    candidate = json.loads(json.dumps(example, ensure_ascii=False))
                    candidate[field][0] = "错误字符串"
                    input_path = directory_path / f"invalid-object-{index}.json"
                    input_path.write_text(json.dumps(candidate, ensure_ascii=False), encoding="utf-8")
                    result = subprocess.run(
                        [
                            sys.executable,
                            str(self.tax / "scripts/generate_report_html.py"),
                            str(input_path),
                            "--validate-only",
                            "--metrics-json",
                            str(metrics),
                        ],
                        capture_output=True,
                        text=True,
                    )
                    self.assertNotEqual(result.returncode, 0)
                    self.assertIn(expected, result.stderr)

    def test_generator_rejects_generic_roadmap_completion_statuses(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            metrics = self.generate_metrics(directory_path)
            example = json.loads(
                (self.tax / "references/report-data.example.json").read_text(encoding="utf-8")
            )
            for index, status in enumerate(("完成", "已完成", "通过", "达标", "待完成")):
                with self.subTest(status=status):
                    candidate = json.loads(json.dumps(example, ensure_ascii=False))
                    candidate["roadmap"][0]["completion"] = status
                    input_path = directory_path / f"generic-completion-{index}.json"
                    input_path.write_text(json.dumps(candidate, ensure_ascii=False), encoding="utf-8")
                    result = subprocess.run(
                        [
                            sys.executable,
                            str(self.tax / "scripts/generate_report_html.py"),
                            str(input_path),
                            "--validate-only",
                            "--metrics-json",
                            str(metrics),
                        ],
                        capture_output=True,
                        text=True,
                    )
                    self.assertNotEqual(result.returncode, 0)
                    self.assertIn(
                        "roadmap completion must describe a verifiable outcome",
                        result.stderr,
                    )

    def test_optional_kpi_note_and_empty_supplemental_calculations_are_valid(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            metrics = self.generate_metrics(directory_path)
            candidate = json.loads(
                (self.tax / "references/report-data.example.json").read_text(encoding="utf-8")
            )
            candidate["financial_overview"]["kpis"][0].pop("note")
            candidate["calculations"] = []
            input_path = directory_path / "optional-fields.json"
            output_path = directory_path / "optional-fields.html"
            input_path.write_text(json.dumps(candidate, ensure_ascii=False), encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(self.tax / "scripts/generate_report_html.py"),
                    str(input_path),
                    str(output_path),
                    "--metrics-json",
                    str(metrics),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            html = output_path.read_text(encoding="utf-8")
            self.assertIn("2025年研发费用率", html)
            self.assertIn("final-indicator-table", html)
            self.assertNotIn("None", html)

    def test_missing_official_policy_generates_an_explicit_draft(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            metrics = self.generate_metrics(directory_path)
            candidate = json.loads(
                (self.tax / "references/report-data.example.json").read_text(encoding="utf-8")
            )
            candidate["policies"] = []
            input_path = directory_path / "draft.json"
            output_path = directory_path / "draft.html"
            input_path.write_text(json.dumps(candidate, ensure_ascii=False), encoding="utf-8")

            subprocess.run(
                [
                    sys.executable,
                    str(self.tax / "scripts/generate_report_html.py"),
                    str(input_path),
                    str(output_path),
                    "--metrics-json",
                    str(metrics),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            html = output_path.read_text(encoding="utf-8")
            self.assertEqual(html.count('<section class="page'), 17)
            self.assertIn("草稿：政策原文未核验", html)
            self.assertIn("本报告为草稿，不得作为正式税务结论使用", html)

    def test_placeholder_policy_is_rejected_instead_of_presented_as_verified(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            metrics = self.generate_metrics(directory_path)
            candidate = json.loads(
                (self.tax / "references/report-data.example.json").read_text(encoding="utf-8")
            )
            candidate["policies"] = [{
                "name": "研发费用加计扣除政策",
                "issuer": "现行状态待核验",
                "date": "待核验",
                "url": "https://example.invalid/pending",
            }]
            input_path = directory_path / "placeholder-policy.json"
            input_path.write_text(json.dumps(candidate, ensure_ascii=False), encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    str(self.tax / "scripts/generate_report_html.py"),
                    str(input_path),
                    "--validate-only",
                    "--metrics-json",
                    str(metrics),
                ],
                capture_output=True,
                text=True,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("placeholder policy is not allowed", result.stderr)

    def test_explicit_word_keeps_tax_html_content_tables_and_draft_status(self):
        from docx import Document
        from docx.oxml.ns import qn
        from lxml import html
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            metrics = self.generate_metrics(root)
            data = json.loads((self.tax / "references/report-data.example.json").read_text(encoding="utf-8"))
            data["policies"] = []
            data["missing_documents"].append("GC-QA 银行流水尚未取得，保持待核验")
            report_input = root / "report-data.json"
            report_input.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            source = root / "report.html"
            subprocess.run([
                sys.executable, str(self.tax / "scripts/generate_report_html.py"),
                str(report_input), str(source), "--metrics-json", str(metrics),
            ], check=True, capture_output=True, text=True)
            original_html = source.read_bytes()
            original_metrics = metrics.read_bytes()
            output = root / "report.docx"
            operations = {
                operation["id"]: operation for operation in json.loads(
                    (self.skills / "client-runtime-operations.json").read_text(encoding="utf-8")
                )["operations"]
            }
            for operation_id, parameters in [
                ("evidence-ledger.create-docx-from-file", {"input": source, "output": output}),
                ("evidence-ledger.apply-office-branding-gold", {"artifact": output}),
            ]:
                operation = operations[operation_id]
                argv = [item["literal"] if "literal" in item else str(parameters[item["parameter"]])
                        for item in operation["argv"]]
                result = subprocess.run([
                    sys.executable, str(self.skills / operation["script"]), *argv,
                ], check=True, capture_output=True, text=True)
                self.assertEqual(json.loads(result.stdout)["status"], "passed")

            document = Document(output)
            source_body = html.document_fromstring(original_html).find("body")
            self.assertIsNotNone(source_body)
            expected_headers = []
            expected_footers = []
            expected_cover_signatures = []
            furniture = source_body.xpath(
                './/*[(self::footer and contains(concat(" ",normalize-space(@class)," ")," page-footer "))'
                ' or contains(concat(" ",normalize-space(@class)," ")," gongchuang-document-header ")'
                ' or contains(concat(" ",normalize-space(@class)," ")," gongchuang-cover-signature ")]'
            )
            for node in furniture:
                value = "".join(node.itertext()).strip()
                if "gongchuang-cover-signature" in node.get("class", "").split():
                    expected_cover_signatures.append(value)
                elif node.tag == "footer":
                    expected_footers.append(value)
                else:
                    expected_headers.append(value)
                node.drop_tree()
            expected_text = "".join("".join(source_body.itertext()).split())
            actual_text = "".join("".join(
                node.text or "" for node in document.element.body.iter(qn("w:t"))
            ).split())
            self.assertEqual(actual_text, expected_text)
            word_section = document.sections[0]
            self.assertEqual([p.text for p in word_section.header.paragraphs if p.text], list(dict.fromkeys(expected_headers)))
            self.assertEqual([p.text for p in word_section.footer.paragraphs if p.text], list(dict.fromkeys(expected_footers)))
            self.assertEqual([p.text for p in word_section.first_page_footer.paragraphs if p.text], list(dict.fromkeys(expected_cover_signatures)))
            self.assertTrue(word_section.different_first_page_header_footer)
            self.assertIn("草稿：政策原文未核验", actual_text)
            self.assertIn("GC-QA银行流水尚未取得，保持待核验", actual_text)
            expected_tables = [
                [["".join(cell.itertext()).strip() for cell in row if cell.tag in {"th", "td"}]
                 for row in table.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")]
                for table in source_body.xpath(".//table")
            ]
            self.assertEqual(
                [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
                expected_tables,
            )
            chapter_titles = [p for p in document.paragraphs if p.text.startswith("01 ")]
            self.assertEqual([p.text for p in chapter_titles], ["01 执行摘要"])
            self.assertTrue(chapter_titles[0].paragraph_format.keep_with_next)
            self.assertTrue(chapter_titles[0].paragraph_format.keep_together)
            self.assertTrue(all(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader") for table in document.tables))
            profile = json.loads((self.skills / "delivery-contracts.json").read_text())["delivery_profiles"]["manufacturing-tax-risk-report"]
            word_role = next(item for item in profile["required_artifacts"] if item["role"] == "tax_report_docx")
            self.assertEqual(word_role["formats"], ["docx"])
            self.assertTrue(any(rule.get("variant") == "gold" and "tax_report_docx" in rule["artifact_roles"]
                                for rule in profile["branding_contracts"]))
            for section in profile["required_sections"]:
                self.assertIn(section, actual_text)
            for table in profile["required_tables"]:
                self.assertIn(table["id"], actual_text)
            verified = subprocess.run([
                sys.executable, str(self.runtime / "scripts/delivery_gate.py"), str(output),
            ], check=True, capture_output=True, text=True)
            self.assertEqual(json.loads(verified.stdout)["status"], "passed")
            gold_assets = {p.read_bytes() for p in (self.runtime / "assets").glob("brand-gold-*.png")}
            with ZipFile(output) as package:
                images = [package.read(name) for name in package.namelist() if name.startswith("word/media/")]
            self.assertTrue(images)
            self.assertTrue(all(image in gold_assets for image in images))
            self.assertEqual(source.read_bytes(), original_html)
            self.assertEqual(metrics.read_bytes(), original_metrics)

    def test_portable_scripts_do_not_contain_machine_paths(self):
        forbidden = ("/Users/", "/Volumes/", ".agents/skills")
        paths = [
            self.tax / "scripts/brand_gold_pdf.py",
            self.tax / "scripts/generate_report_html.py",
            self.tax / "scripts/render_pdf_stdout.js",
            self.tax / "scripts/verify_e2e.py",
            self.runtime / "scripts/brand_config.py",
            self.runtime / "scripts/pdf_two_pass.py",
            self.runtime / "scripts/delivery_gate.py",
        ]
        for path in paths:
            content = path.read_text(encoding="utf-8")
            for marker in forbidden:
                self.assertNotIn(marker, content, f"{path} contains {marker}")

    def test_branding_runtime_imports_from_its_own_directory(self):
        module_path = self.runtime / "scripts/brand_config.py"
        spec = importlib.util.spec_from_file_location("portable_brand_config", module_path)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        config = module.load_config()
        self.assertEqual(config["policy"]["position"], [0.5, 0.5])
        style = module.choose_style(0.5, target="pdf", variant="gold")
        self.assertTrue(Path(style["asset_path"]).is_file())


if __name__ == "__main__":
    unittest.main()
