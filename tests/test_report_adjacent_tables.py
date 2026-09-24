from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def test_distinct_report_tables_have_separate_word_headers(tmp_path):
    script = Path(__file__).resolve().parents[1] / "skills/evidence-ledger/scripts/create_docx_from_text.py"
    spec = spec_from_file_location("report_docx_tables", script)
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    output = tmp_path / "report.docx"
    module.build_document("# 材料体检\n\n| 产品 | 技术 |\n|---|---|\n| PS01 | 补偿 |\n\n| IP | 状态 |\n|---|---|\n| IP01 | 授权 |\n", output)
    doc = Document(output)
    assert [[c.text for c in t.rows[0].cells] for t in doc.tables] == [["产品", "技术"], ["IP", "状态"]]
    assert doc.tables[0]._tbl.getnext().tag == qn("w:p")
    assert doc.tables[0]._tbl.getnext().getnext() is doc.tables[1]._tbl
