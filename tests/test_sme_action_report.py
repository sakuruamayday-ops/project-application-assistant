import copy
import importlib.util
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills/project-feasibility/scripts"
spec = importlib.util.spec_from_file_location("sme_action_report", SCRIPTS / "sme_action_report.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def fixture():
    return {
        "enterprise": "合成测试企业", "project_id": "specialized-sme",
        "report_date": "2026-09-28", "suggested_year": "2027",
        "report_sections": {
            key: {"paragraphs": [f"{key}的实际说明"], "rows": [[f"{key}-{i}" for i in range(len(headers))]] if headers else []}
            for key, _, headers in module.SECTIONS
        },
    }


def build(data):
    return module.build_action_document(
        ROOT / "skills/project-feasibility/assets/report-templates/specialized-sme/preassessment.docx", data
    )


def tasks(doc):
    return next(t for t in doc.tables if t.rows[0].cells[0].text == "补强动作")


@pytest.mark.parametrize("assets,expected", [(None, []), ("0", []),
    ("5999999.99", []), ("6000000", ["研发档A"]),
    ("10000000", ["研发档A", "研发档B"])])
def test_research_equipment_half_assets_boundary(assets, expected):
    tiers = [{"name": "研发档B", "equipment_threshold_yuan": "5000000", "source": "合成政策B"},
             {"name": "研发档A", "equipment_threshold_yuan": "3000000", "source": "合成政策A"}]
    assert module.screen_research_equipment(assets, tiers)["recommended"] == expected


@pytest.mark.parametrize("assets", ["-1", "NaN", "Infinity", "错误"])
def test_research_equipment_rejects_invalid_amounts(assets):
    with pytest.raises(ValueError, match="非负有限数"):
        module.screen_research_equipment(assets, [])


def test_research_screen_reaches_renderer_without_leaking_estimate():
    data = fixture()
    data["research_equipment_screen"] = {
        "fixed_assets_yuan": "6000000", "existing_basis": "已有研发团队及产品",
        "tiers": [{"name": "合成研发中心", "equipment_threshold_yuan": "3000000", "source": "测试政策"}],
    }
    before = copy.deepcopy(data)
    doc = build(data)
    text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "建议申报方向：合成研发中心" in text
    assert "6000000" not in text and "3000000" not in text and "50%" not in text
    assert data == before
    assert [c.text for c in tasks(doc).rows[1].cells] == data["report_sections"]["tasks"]["rows"][0]


def test_research_screen_never_uses_total_assets():
    data = fixture()
    data["research_equipment_screen"] = {
        "total_assets_yuan": "100000000", "existing_basis": "已有研发团队及产品",
        "tiers": [{"name": "合成研发中心", "equipment_threshold_yuan": "3000000", "source": "测试政策"}],
    }
    rows = module.prepare_sections(data)["next_tasks"]["rows"]
    assert rows[-1][1] == "优先培育"
    assert "合成研发中心" not in rows[-1][-1]


@pytest.mark.parametrize("project", ["specialized-sme", "little-giant"])
def test_fixed_structure_and_preserved_tasks(project):
    data = fixture()
    data["project_id"] = project
    doc = build(data)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == [heading for heading, _ in module.GROUPS]
    assert [c.text for c in tasks(doc).rows[1].cells] == data["report_sections"]["tasks"]["rows"][0]
    assert "5.1 补强任务表" in [p.text for p in doc.paragraphs]


def test_edit_one_section_does_not_reset_other_content():
    before = fixture()
    after = copy.deepcopy(before)
    after["report_sections"]["product"]["paragraphs"] = ["用户修改后的产品说明"]
    doc = build(after)
    assert "用户修改后的产品说明" in [p.text for p in doc.paragraphs]
    assert [c.text for c in tasks(doc).rows[1].cells] == before["report_sections"]["tasks"]["rows"][0]
    assert doc.settings.element.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}documentProtection") is None


def test_missing_product_has_explanation_not_empty_template_rows():
    data = fixture()
    data["report_sections"]["product"] = {"paragraphs": ["未提供产品清单，先核实际销售对象。"], "rows": []}
    assert "未提供产品清单，先核实际销售对象。" in [p.text for p in build(data).paragraphs]


@pytest.mark.parametrize("section", ["product", "peers"])
def test_empty_business_section_is_not_a_complete_report(section):
    data = fixture()
    data["report_sections"][section] = {"paragraphs": [], "rows": []}
    with pytest.raises(ValueError, match=section):
        build(data)


@pytest.mark.parametrize("source,expected", [
    ("拥有三项技术优势，评分为70分。", "拥有三项技术优势。"),
    ("评分为70分，拥有三项技术优势。", "拥有三项技术优势。"),
    ("综合评价：70 / 100。拥有三项技术优势。", "拥有三项技术优势。"),
    ("对应评分项，可形成加分优势。", "对应评价指标，可形成加分优势。"),
    ("处理时间50分钟，合格样本70 / 100。", "处理时间50分钟，合格样本70 / 100。"),
    ("收入10,000万元，研发人员60名。", "收入10,000万元，研发人员60名。"),
])
def test_score_filter_preserves_non_score_facts(source, expected):
    assert module.public_text(source) == expected


@pytest.mark.parametrize("label", ["前置身份", "科技和创新型中小企业", "取得科技和创新型中小企业称号", "创新型中小企业（省级科技型中小企业）"])
def test_provincial_prerequisite_aliases_are_hidden_without_changing_input(label):
    data = fixture()
    data["report_sections"]["tasks"]["rows"].append([label, "确认称号有效状态", "申报前"])
    data["report_sections"]["conclusion"]["paragraphs"] = [
        f"研发费用126万元，{label}尚待确认。主导产品收入1260万元。"
    ]
    before = copy.deepcopy(data)
    doc = build(data)
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert label not in text
    assert "研发费用126万元。主导产品收入1260万元。" in text
    assert data == before


def test_little_giant_provincial_title_remains_visible_in_combined_report():
    text = "小巨人要求已获得省级专精特新中小企业称号，企业尚未取得。"
    assert module.public_text(text, "specialized-sme") == text
    assert module.public_text(text, "little-giant") == text


@pytest.mark.parametrize("section", ["tasks", "soft"])
def test_qualitative_scoring_task_is_not_removed(section):
    data = fixture()
    row = ["PCT国际布局", "对应评分项，可形成加分优势。", "申报前"]
    if section == "soft":
        row = ["PCT国际布局", "当前状态待核", row[1], "申报前"]
    data["report_sections"][section]["rows"] = [row]
    before = copy.deepcopy(data)
    doc = build(data)
    contents = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "PCT国际布局" in contents
    assert "对应评价指标，可形成加分优势。" in contents
    assert data == before


@pytest.mark.parametrize("section", ["product", "peers"])
def test_filtered_only_section_is_not_complete(section):
    data = fixture()
    data["report_sections"][section] = {"paragraphs": ["综合评价：70 / 100。"], "rows": []}
    with pytest.raises(ValueError, match=section):
        build(data)


@pytest.mark.parametrize("project", ["specialized-sme", "little-giant"])
def test_exported_docx_retains_facts_and_task_after_reopening(tmp_path, project):
    from docx import Document

    data = fixture()
    data["project_id"] = project
    data["report_sections"]["conclusion"]["paragraphs"] = ["拥有三项技术优势，评分为70分。"]
    data["report_sections"]["tasks"]["rows"] = [["PCT国际布局", "对应评分项，可形成加分优势。", "申报前"]]
    path = tmp_path / f"{project}.docx"
    build(data).save(path)
    reopened = Document(path)
    assert "拥有三项技术优势。" in [p.text for p in reopened.paragraphs]
    assert [c.text for c in tasks(reopened).rows[1].cells] == [
        "PCT国际布局", "对应评价指标，可形成加分优势。", "申报前",
    ]
    assert "70分" not in "\n".join(p.text for p in reopened.paragraphs)


@pytest.mark.parametrize("mutation", ["wrong-project", "missing-section", "wrong-columns", "empty-tasks"])
def test_invalid_input_has_direct_error(mutation):
    data = fixture()
    if mutation == "wrong-project":
        data["project_id"] = "high-tech-enterprise"
    elif mutation == "missing-section":
        del data["report_sections"]["finance"]
    elif mutation == "wrong-columns":
        data["report_sections"]["tasks"]["rows"] = [["x"]]
    else:
        data["report_sections"]["tasks"]["rows"] = []
    with pytest.raises(ValueError):
        build(data)


def test_runtime_manifest_includes_renderer_dependency():
    import json
    operations = json.loads((ROOT / "skills/client-runtime-operations.json").read_text())["operations"]
    operation = next(o for o in operations if o["id"] == "project-feasibility.generate-report")
    assert "project-feasibility/scripts/sme_action_report.py" in operation["files"]


@pytest.mark.parametrize("project", ["specialized-sme", "little-giant"])
def test_scores_and_unavailable_score_explanations_not_exported(project):
    data = fixture()
    data["project_id"] = project
    data["report_sections"]["conclusion"]["paragraphs"] = ["研发基础保留。分数由系统自动评，获取不到。"]
    data["report_sections"]["acceptance"]["rows"].append(
        ["发展质量", "当年平台得分≥60分", "当前平台分未取得", "明年模拟填报"]
    )
    doc = build(data)
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "研发基础保留。" in text
    assert not module.SCORE_TEXT.search(text)
    assert "获取不到" not in text
    assert data["report_sections"]["acceptance"]["rows"][-1][0] == "发展质量"


def test_all_reference_content_blocks_exist_in_order():
    doc = build(fixture())
    expected = [heading for _, heading, _ in module.SECTIONS if heading]
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2"] == expected


def test_action_values_and_peer_comparison_columns():
    doc = build(fixture())
    assert [c.text for c in tasks(doc).rows[0].cells] == ["补强动作", "补强价值", "建议节点"]
    peer = next(t for t in doc.tables if t.cell(0, 0).text == "对标企业或产品体系")
    assert len(peer.columns) == 4
    assert peer.cell(0, 2).text == "与本企业的异同"


@pytest.mark.parametrize("project", ["specialized-sme", "little-giant"])
def test_qualitative_indicator_benefits_survive_export(project):
    data = fixture()
    data["project_id"] = project
    data["report_sections"]["tasks"]["rows"] = [["PCT国际布局", "对应国际拓展正向指标，可作为加分提升方向。", "按适用期限"]]
    data["report_sections"]["soft"]["rows"] = [["质量管理能力", "等级待核", "对应质量控制，可形成加分优势。", "按实际等级取用"]]
    doc = build(data)
    text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "加分提升方向" in text and "加分优势" in text
    assert not module.SCORE_TEXT.search(text)


@pytest.mark.parametrize("project", ["specialized-sme", "little-giant"])
def test_project_specific_prerequisite_and_hidden_sections(project):
    data = fixture()
    data["project_id"] = project
    data["report_sections"]["conditional"] = {"paragraphs": ["条件触发旧输入"], "rows": []}
    data["report_sections"]["current_tasks"]["rows"].append(
        ["前置称号", "近期", "旧前置资格条件", "企业状态待核", "核对", "核认定文件与有效状态"])
    data["report_sections"]["acceptance"]["rows"].append(["资格与专注", "条件", "现状", "动作"])
    doc = build(data)
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "条件触发" not in text and "资格与专注" not in text
    if project == "specialized-sme":
        assert "前置称号" not in text
    else:
        assert "前置称号" in text and "已获认定的省级专精特新中小企业" in text
        assert "企业状态待核" in text
    assert data["report_sections"]["current_tasks"]["rows"][-1][2] == "旧前置资格条件"
