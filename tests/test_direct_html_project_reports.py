import copy
import importlib
import json
import subprocess
import sys
from pathlib import Path

import pytest
from lxml import html


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills/project-feasibility/scripts"
PROJECTS = ("high-tech-enterprise", "specialized-sme", "little-giant")


@pytest.fixture
def report_modules(monkeypatch):
    monkeypatch.syspath_prepend(str(SCRIPTS))
    return importlib.import_module("generate_report"), importlib.import_module("fill_report_template")


def case_fixture(tmp_path, project):
    source = tmp_path / "synthetic-source.txt"
    source.write_text("合成测试企业研发智能控制系统，已提供研发费用126万元。", encoding="utf-8")
    return {
        "project_id": project, "enterprise": "合成测试企业", "project_object": "智能控制系统",
        "report_date": "2026-10-12", "suggested_year": "2027年", "deadline": "以当期通知为准",
        "conclusion": "有条件申报", "conclusion_basis": "已提供技术资料，财务及权利状态待核。",
        "primary_gap": "收入及研发费用口径待核", "next_action": "核对财务及产品技术台账",
        "materials": [{"path": str(source), "role": "合成技术材料", "anchors": ["合成测试企业", "智能控制系统"],
                       "source_type": "knowledge-base"}],
        "policies": [{"title": "合成申报年度通知", "locator": "https://example.gov.cn/test-policy",
                      "kind": "annual-notice", "status": "合成测试来源"}],
    }


def generate(module, tmp_path, data, report_type, filename="report.html"):
    input_path = tmp_path / "facts.json"
    input_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / filename
    result = module.generate_report(input_path, output, report_type)
    return result, html.fromstring(output.read_text(encoding="utf-8"))


@pytest.mark.parametrize("project", PROJECTS)
@pytest.mark.parametrize("report_type", ["preassessment", "feasibility"])
def test_report_generation_never_reads_a_word_master(tmp_path, monkeypatch, report_modules, project, report_type):
    generator, filler = report_modules
    data = case_fixture(tmp_path, project)

    def reject_word(*args, **kwargs):
        raise AssertionError("report generation must not read or serialize a Word master")

    monkeypatch.setattr(filler, "Document", reject_word)
    result, document = generate(generator, tmp_path, data, report_type)
    assert result["status"] == "pass"
    assert Path(result["template_path"]).suffix == ".html"
    text = document.find("body").text_content()
    assert data["enterprise"] in text and data["project_object"] in text
    assert "数据来源" in text
    assert len(document.xpath("//table")) >= 5
    assert not list(tmp_path.glob("*.docx")) and not list(tmp_path.glob("*.pdf"))
    assert document.xpath('//a[@href="https://example.gov.cn/test-policy"]')
    assert not any(marker in text for marker in ("使用说明", "判定规则", "项目路由", "专项要求"))
    assert document.xpath('//th[.//text()="官方来源"]//span[contains(@style,"#FFFFFF")]')


@pytest.mark.parametrize("project", PROJECTS)
@pytest.mark.parametrize("report_type", ["preassessment", "feasibility"])
def test_analysis_report_does_not_enter_the_application_word_path(tmp_path, report_modules, project, report_type):
    generator, _ = report_modules
    with pytest.raises(ValueError, match="HTML|html"):
        generate(generator, tmp_path, case_fixture(tmp_path, project), report_type, "report.docx")
    assert not (tmp_path / "report.docx").exists()
    assert not (tmp_path / "report.completion.json").exists()


@pytest.mark.parametrize("project", ["specialized-sme", "little-giant"])
def test_action_report_keeps_fixed_sections_and_user_edits_without_word(tmp_path, monkeypatch, report_modules, project):
    generator, filler = report_modules
    action = importlib.import_module("sme_action_report")
    data = case_fixture(tmp_path, project)
    data["report_layout"] = "sme-action"
    data["report_sections"] = {
        key: {"paragraphs": [f"{key}合成说明"],
              "rows": [[f"{key}-{i}" for i in range(len(headers))]] if headers else []}
        for key, _, headers in action.SECTIONS
    }
    data["report_sections"]["tasks"]["rows"] = [["PCT国际布局", "对应评价指标，可形成加分优势。", "申报前"]]
    data["report_sections"]["product"]["paragraphs"] = ["用户修订后的智能控制系统产品说明"]
    before = copy.deepcopy(data)

    def reject_word(*args, **kwargs):
        raise AssertionError("action reports must not read a Word master")

    monkeypatch.setattr(filler, "Document", reject_word)
    monkeypatch.setattr(action, "Document", reject_word)
    result, document = generate(generator, tmp_path, data, "preassessment")
    headings = [node.text_content() for node in document.xpath("//body/h1")]
    assert headings[1:6] == [title for title, _ in action.GROUPS]
    assert "5.1 补强任务表" in document.text_content()
    assert "用户修订后的智能控制系统产品说明" in document.text_content()
    assert "PCT国际布局" in document.text_content() and "可形成加分优势" in document.text_content()
    assert result["delivery_profile"] == "sme-action-report"
    assert data == before


def test_uploaded_word_is_read_but_no_word_report_master_is_used(tmp_path, monkeypatch, report_modules):
    from docx import Document
    generator, filler = report_modules
    data = case_fixture(tmp_path, "high-tech-enterprise")
    source = tmp_path / "uploaded-application.docx"
    document = Document()
    document.add_paragraph("合成测试企业研发智能控制系统，研发费用126万元。")
    document.save(source)
    data["materials"][0]["path"] = str(source)
    calls = []

    def read_uploaded_file(path):
        assert isinstance(path, Path) and path == source
        calls.append(path)
        return Document(path)

    monkeypatch.setattr(filler, "Document", read_uploaded_file)
    generate(generator, tmp_path, data, "feasibility")
    assert calls == [source]


@pytest.mark.parametrize("report_type", ["preassessment", "feasibility"])
def test_html_preserves_complete_rd_rows_numbers_and_bracketed_source_values(tmp_path, report_modules, report_type):
    generator, _ = report_modules
    data = case_fixture(tmp_path, "high-tech-enterprise")
    data["core_associations"] = [[f"RD{i:02}", "技术链", "状态待核", "按［2026］36号原始资料列示",
                                  "PS01收入126万元，未按RD重复计收入", "合成技术台账"] for i in range(1, 18)]
    _, document = generate(generator, tmp_path, data, report_type)
    matrix = next(table for table in document.xpath("//table")
                  if table.xpath(".//th")[0].text_content() == "研发项目")
    rows = matrix.xpath("./tbody/tr")[1:]
    assert len(rows) == 17
    for actual, expected in zip(rows, data["core_associations"], strict=True):
        assert [cell.text_content() for cell in actual.xpath("./td")] == expected


@pytest.mark.parametrize("columns,rows", [
    (["序号", "评分维度", "满分", "预计", "评分规则", "证据", "补强"], [["错列"]])])
def test_invalid_html_table_keeps_final_path_available(tmp_path, report_modules, columns, rows):
    generator, _ = report_modules
    data = case_fixture(tmp_path, "high-tech-enterprise")
    data["table_content"] = [{"columns": columns, "rows": rows}]
    with pytest.raises(ValueError, match="表头|列数"):
        generate(generator, tmp_path, data, "preassessment")
    assert not (tmp_path / "report.html").exists()
    assert not (tmp_path / "report.completion.json").exists()
    data.pop("table_content")
    generate(generator, tmp_path, data, "preassessment")


def test_html_values_are_text_not_executable_markup_and_existing_file_is_retained(tmp_path, report_modules):
    generator, _ = report_modules
    data = case_fixture(tmp_path, "high-tech-enterprise")
    data["primary_gap"] = '<script>alert("test")</script> & 客户原文'
    _, document = generate(generator, tmp_path, data, "preassessment")
    assert not document.xpath("//script")
    assert data["primary_gap"] in document.text_content()
    before = (tmp_path / "report.html").read_bytes()
    with pytest.raises(FileExistsError):
        generate(generator, tmp_path, data, "preassessment")
    assert (tmp_path / "report.html").read_bytes() == before


@pytest.mark.parametrize("invalid", ["word-report", "table-width"])
def test_cli_returns_one_explicit_json_failure_without_creating_an_artifact(tmp_path, invalid):
    data = case_fixture(tmp_path, "high-tech-enterprise")
    if invalid == "table-width":
        data["table_content"] = [{"columns": ["指标"], "rows": [["内容", "错列"]]}]
    input_path = tmp_path / "facts.json"
    input_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / ("report.docx" if invalid == "word-report" else "report.html")
    result = subprocess.run([sys.executable, str(SCRIPTS / "generate_report.py"), "--input", str(input_path),
                             "--output", str(output), "--report-type", "preassessment"], capture_output=True, text=True)
    payload = json.loads(result.stdout)
    assert result.returncode == 2
    assert payload["schema_version"] == "gongchuang-project-report-operation/v1"
    assert payload["status"] == "fail" and payload["errors"]
    assert "HTML" in payload["errors"][0] if invalid == "word-report" else "列数" in payload["errors"][0]
    if invalid == "table-width":
        assert payload["input_rejection"] == {"code": "report-table-shape", "fields": ["table_content.rows"]}
    assert not output.exists()
    operation = next(op for op in json.loads((ROOT / "skills/client-runtime-operations.json").read_text())["operations"]
                     if op["id"] == "project-feasibility.generate-report")
    assert 2 in operation["result_exit_codes"] and 2 not in operation["passing_exit_codes"]


def test_bom_missing_sources_and_extra_table_deliver_a_visible_draft(tmp_path, report_modules):
    generator, _ = report_modules
    data = {"project_id": "high-tech-enterprise", "enterprise": "合成测试企业",
            "table_content": {"columns": ["补充业务事实"], "rows": [["独立保留的事实"]]}}
    input_path = tmp_path / "bom-facts.json"
    input_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8-sig")
    output = tmp_path / "draft.html"
    result = generator.generate_report(input_path, output, "preassessment")
    document = html.fromstring(output.read_text(encoding="utf-8"))
    assert result["status"] == "pass" and result["report_stage"] == "draft"
    assert result["source_count"] == 0 and result["warnings"]
    assert "初稿说明与待补事项" in document.text_content()
    assert "独立保留的事实" in document.text_content()
    assert "尚未提供可核验的客户资料" in document.text_content()
    assert "申报核心对象尚未明确" in document.text_content()
    assert "project_object" not in document.text_content()
    assert "暂无法判断" in document.text_content()
    assert "materials" not in data and "_input_warnings" not in data


def test_evidence_and_table_header_aliases_preserve_the_master(tmp_path, report_modules):
    generator, filler = report_modules
    data = case_fixture(tmp_path, "high-tech-enterprise")
    data["conditions"] = [{"item": "独立输入的条件", "state": "待补强"}]
    data["table_content"] = {"columns": ["序号", "评分维度", "满分值", "预计得分", "评分依据", "证据", "补强建议"],
                             "rows": [["1", "合成维度", "10", "待核", "官方规则", "合成来源", "补足资料"]]}
    normalized = filler.validate_case_fixture(data)
    assert normalized["conditions"][0]["state"] == "存在差距"
    assert data["conditions"][0]["state"] == "待补强"
    result, document = generate(generator, tmp_path, data, "feasibility")
    assert "合成维度" in document.text_content()
    assert "补充分析" not in document.text_content()
    assert not result["warnings"]
