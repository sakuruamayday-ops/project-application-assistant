import importlib.util
import json
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills/high-tech-enterprise-application-drafting"
SCRIPT = SKILL / "scripts/fill_application_fields.py"
TEMPLATE = SKILL / "assets/高新技术企业认定申请书空白模板.docx"
MODULE_SPEC = importlib.util.spec_from_file_location("expand_application_fields_test", SKILL / "scripts/expand_rd_ps_tables.py")
EXPAND = importlib.util.module_from_spec(MODULE_SPEC)
MODULE_SPEC.loader.exec_module(EXPAND)


def run_fill(tmp_path, data, template=TEMPLATE):
    if template == TEMPLATE:
        document = Document(template)
        count = max(1, len(data.get("rd_financials", {})))
        EXPAND.resize_kind(document, "rd", count, "RD")
        EXPAND.resize_kind(document, "ps", 1, "PS")
        template = tmp_path / "expanded.docx"
        document.save(template)
    spec = tmp_path / "input.json"
    spec.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "result.docx"
    result = subprocess.run([sys.executable, str(SCRIPT), str(template), str(spec), str(output)], capture_output=True, text=True)
    return result, output


def physical(table, row, col):
    return "".join(table._tbl.tr_lst[row].tc_lst[col].xpath(".//w:t/text()"))


def test_merged_registration_labels_and_financial_totals(tmp_path):
    result, output = run_fill(tmp_path, {
        "registration": {"企业名称": "GC-QA测试", "注册时间": "2020年6月1日"},
        "summary": {"Ⅰ类": 2, "Ⅱ类": None, "职工总数": 60, "科技人员数": 12},
        "annual_financials": [
            {"year": 2023, "net_assets": 600, "sales": 1200, "profit": 96, "rd": 84, "domestic_rd": 84},
            {"year": 2024, "net_assets": 690, "sales": 1500, "profit": 120, "rd": 105, "domestic_rd": 105},
            {"year": 2025, "net_assets": 793.5, "sales": 1800, "profit": 144, "rd": 126, "domestic_rd": 126},
        ],
        "rd_financials": {"RD01": [84, 55, 0], "RD02": [0, 50, 76], "RD03": [0, 0, 50]},
        "rd_external": {"RD01": [0, 0, 0], "RD02": [0, 0, 0], "RD03": [0, 0, 0]},
        "personnel": {"总   数（人）": [60, 12]},
    })
    assert result.returncode == 0, result.stderr
    tables = Document(output).tables
    assert [physical(tables[0], 0, c) for c in range(4)] == ["企业名称", "GC-QA测试", "注册时间", "2020年6月1日"]
    assert physical(tables[1], 1, 3) == "Ⅱ类"
    assert physical(tables[1], 2, 3) == "科技人员数"
    assert physical(tables[1], 7, 1) == "315"
    assert physical(tables[1], 7, 4) == "315"
    assert physical(tables[5], 3, 3) == "139"
    assert [physical(tables[5], r, 6) for r in [3, 4, 5]] == ["84", "55", "0"]
    assert physical(tables[4], 2, 1) == "60"
    assert physical(tables[4], 2, 2) == "12"
    # All table grid and merge properties survive value writes.
    original = Document(tmp_path / "expanded.docx")
    expected = original.tables[:8] + [original.tables[8]] * 3 + original.tables[9:]
    assert len(tables) == len(expected)
    for before, after in zip(expected, tables):
        assert before._tbl.tblGrid.xml == after._tbl.tblGrid.xml
        assert [cell.tcPr.xml for row in before._tbl.tr_lst for cell in row.tc_lst] == [cell.tcPr.xml for row in after._tbl.tr_lst for cell in row.tc_lst]
    assert [physical(tables[i], 11, 6) for i in [8, 9, 10]] == ["84", "105", "126"]
    assert [physical(tables[i], 1, 6) for i in [8, 9, 10]] == ["84", "105", "126"]
    document = Document(output)
    for i, year in zip([8, 9, 10], [2023, 2024, 2025]):
        assert not tables[i]._tbl.xpath("./w:tblPr/w:tblpPr")
        caption = tables[i]._tbl.getprevious()
        assert "".join(caption.xpath(".//w:t/text()")).startswith(str(year))
        assert caption.xpath("./w:pPr/w:keepNext")
    assert not document.element.xpath("./w:body/w:p//w:br[@w:type='page']")
    for title in ["填  报  说  明", "三、人力资源情况表"]:
        paragraph = next(p for p in document.paragraphs if p.text == title)
        assert paragraph.paragraph_format.page_break_before
    assert "企业名称：GC-QA测试" in "".join(Document(output).element.xpath("//w:body/w:p//w:t/text()"))


def test_missing_cost_is_not_zero(tmp_path):
    result, output = run_fill(tmp_path, {"rd_financials": {"RD01": [84, None, 0]}})
    assert result.returncode == 0, result.stderr
    assert physical(Document(output).tables[5], 3, 3) == ""


@pytest.mark.parametrize("data", [
    {"registration": {"不存在的字段": "value"}},
    {"rd_financials": {"RD99": [1, 2, 3]}},
    {"rd_financials": {"RD01": [1, -2, 3]}},
    {"ps": [{}, {}]},
])
def test_bad_input_does_not_write_partial_file(tmp_path, data):
    result, output = run_fill(tmp_path, data)
    assert result.returncode != 0
    assert not output.exists()


def test_existing_output_is_preserved(tmp_path):
    output = tmp_path / "result.docx"
    output.write_bytes(b"existing")
    result, _ = run_fill(tmp_path, {})
    assert result.returncode != 0
    assert output.read_bytes() == b"existing"


def test_seventeen_rd_projects_split_without_losing_totals(tmp_path):
    costs = {f"RD{i:02d}": [i, i * 2, i * 3] for i in range(1, 18)}
    result, output = run_fill(tmp_path, {
        "annual_financials": [{"year": 2023 + n, "rd": 153 * (n + 1)} for n in range(3)],
        "rd_financials": costs,
    })
    assert result.returncode == 0, result.stderr
    document = Document(output)
    expenses = [table for table in document.tables if "研究开发费用（内、外部）小计" in table._tbl.xml]
    assert len(expenses) == 12
    for year in range(3):
        tables = expenses[year * 4:year * 4 + 4]
        headers = [physical(table, 0, col) for table in tables for col in range(1, 6)]
        assert [value for value in headers if value] == list(costs)
        assert sum(int(physical(table, 11, 6)) for table in tables) == 153 * (year + 1)
        assert physical(tables[-1], 11, 3) == ""
        assert all(table._tbl.tblGrid.xml == tables[0]._tbl.tblGrid.xml for table in tables)
    assert "RD16至RD17，合计为本表项目小计" in "".join(document.element.xpath("//w:t/text()"))


@pytest.mark.parametrize("external,domestic,annual", [([2, 0, 0], [0, 0, 0], 1),
                                                           ([1, 0, 0], [2, 0, 0], 1),
                                                           ([0, 0, 0], [0, 0, 0], 2)])
def test_inconsistent_expense_relationships_do_not_generate(tmp_path, external, domestic, annual):
    result, output = run_fill(tmp_path, {
        "annual_financials": [{"year": 2023 + n, "rd": annual if n == 0 else 0} for n in range(3)],
        "rd_financials": {"RD01": [1, 0, 0]},
        "rd_external": {"RD01": external}, "rd_external_domestic": {"RD01": domestic},
    })
    assert result.returncode != 0
    assert not output.exists()


def test_ps_innovation_and_ip_are_written_to_named_fields(tmp_path):
    result, output = run_fill(tmp_path, {
        "ps": [{"产品（服务）名称": "GC-QA检测装备", "上年度销售收入（万元）": 1260,
                "知识产权编号": "IP01", "关键技术及主要技术指标（限400字）": "夹持重复定位误差不大于0.02毫米。"}],
        "innovation": {"科技成果转化情况（限400字）": "三个年度分别完成2、3、3项转化。"},
        "ip": [{"知识产权名称": "GC-QA合成授权", "类别": "发明专利", "授权号": None}],
        "ip_counts": {"发明专利": 1},
    })
    assert result.returncode == 0, result.stderr
    content = "".join(Document(output).element.xpath("//w:t/text()"))
    for value in ["GC-QA检测装备", "1260", "IP01", "GC-QA合成授权", "三个年度分别完成2、3、3项转化。"]:
        assert value in content


def test_short_field_names_and_embedded_cjk_font(tmp_path):
    result, output = run_fill(tmp_path, {"schema_version": 1,
        "innovation": {"科技成果转化情况": "仅采用已确认的转化事实。"}})
    assert result.returncode == 0, result.stderr
    with ZipFile(output) as package:
        assert "word/fonts/gongchuang-noto-sans-sc.odttf" in package.namelist()
        assert b"embedRegular" in package.read("word/fontTable.xml")
        assert b"Noto Sans SC" in package.read("word/document.xml")
    assert "仅采用已确认的转化事实。" in "".join(Document(output).element.xpath("//w:t/text()"))


@pytest.mark.parametrize("data", [
    {"schema_version": 2}, {"schema_version": True},
    {"innovation": {"科技成果转化情况": "甲", "科技成果转化情况（限400字）": "乙"}},
    {"innovation": {"科技成果转化情况": "甲" * 401}},
])
def test_bad_version_duplicate_or_overlong_field_is_rejected(tmp_path, data):
    result, output = run_fill(tmp_path, data)
    assert result.returncode != 0
    assert not output.exists()
