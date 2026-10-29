"""Office branding retains validation when the PDF native stack is unavailable."""
from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import sys


ROOT = Path(__file__).resolve().parents[1]


def test_office_branding_and_validation_do_not_load_pdf_native_stack(tmp_path):
    probe = r'''
import importlib.abc, json, sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
from docx import Document
from openpyxl import Workbook

class RejectPdf(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname == "pymupdf" or fullname.startswith("pymupdf."):
            raise AssertionError("Office operation attempted to load native PDF stack")

sys.meta_path.insert(0, RejectPdf())
from office_watermark import apply_office_watermark
from delivery_gate import GateFailure, validate_artifact

destination = Path(sys.argv[2])
docx_path = destination / "office.docx"
doc = Document()
doc.add_paragraph("Synthetic Word content")
doc.save(docx_path)
xlsx_path = destination / "office.xlsx"
workbook = Workbook()
workbook.active["A1"] = "Synthetic spreadsheet content"
workbook.save(xlsx_path)
results = []
for path in (docx_path, xlsx_path):
    apply_office_watermark(path)
    result = validate_artifact(path)
    assert result["status"] == "passed"
    results.append(result["format"])

unbranded = destination / "unbranded.docx"
Document().save(unbranded)
try:
    validate_artifact(unbranded)
except GateFailure:
    pass
else:
    raise AssertionError("Unbranded document passed existing validation")
assert "pymupdf" not in sys.modules
assert not any(name.startswith(("pymupdf.layout", "pymupdf4llm", "onnxruntime")) for name in sys.modules)
print(json.dumps({"formats": results, "pdfImported": False, "mlImported": False, "unbrandedRejected": True}))
'''
    result = subprocess.run(
        [os.environ.get("GONGCHUANG_TEST_PYTHON", sys.executable), "-B", "-I", "-c", probe,
         str(ROOT / "skills" / "_runtime" / "gongchuang-branding" / "scripts"), str(tmp_path)],
        check=True, capture_output=True, text=True, timeout=30,
    )
    assert json.loads(result.stdout) == {
        "formats": ["docx", "xlsx"], "pdfImported": False, "mlImported": False, "unbrandedRejected": True,
    }
