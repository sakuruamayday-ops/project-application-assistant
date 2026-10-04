"""Render SME action reports from explicit sections, without placeholder filling."""
from pathlib import Path
from copy import deepcopy
from decimal import Decimal, InvalidOperation
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


GROUPS = (
    ("一、先说结论", ("conclusion", "policy", "known")),
    ("二、项目申报路径图", ("roadmap", "current_tasks", "next_tasks", "acceptance", "soft")),
    ("三、主导产品怎么定", ("product", "product_boundary", "peers")),
    ("四、主导产品怎么配发明专利", ("ip", "ip_topics", "ip_time")),
    ("五、财务情况简析", ("finance", "metrics", "finance_conclusion", "finance_missing", "tasks")),
)
SECTIONS = (
    ("conclusion", "", ("主导产品首选", "申报主线", "立即要做", "最大风险")),
    ("policy", "现行申报核心条件", ("条件组", "申报条件怎么要求")),
    ("known", "只对照企业已有数据的指标", ("关键指标", "企业现有数据", "判断", "谈单时怎么说")),
    ("roadmap", "", ("近期准备", "年度建设", "申报前")),
    ("current_tasks", "当前年度必须完成的资格与底稿", ("项目", "安排", "现行条件", "已有数据", "差距", "立即动作")),
    ("next_tasks", "下一年度的平台与产品项目", ("项目", "安排", "现行条件", "已有数据", "差距", "建议动作")),
    ("acceptance", "主申报前的内部验收", ("内部验收项", "现行明确条件", "企业当前判断", "申报前验收目标")),
    ("soft", "两年任务包中的软提升项", ("软提升项", "现有情况", "对专精线的作用", "两年动作")),
    ("product", "为什么这个名称更合适", ("口径", "候选名称", "结论", "理由与使用条件")),
    ("product_boundary", "产品、收入和专利必须用同一个边界", ("层级", "本企业建议口径", "取证方法")),
    ("peers", "同行与行业对照", ("对标企业或产品体系", "已核验产品与能力", "与本企业的异同", "对本企业的启示")),
    ("ip", "发明的建议分组", ("技术簇", "现有基础", "需要解决的产品问题")),
    ("ip_topics", "建议专利主题和产品价值对应", ("技术簇", "建议发明主题方向", "要证明的产品价值")),
    ("ip_time", "时间要求", ("近期", "建设期", "申报前")),
    ("finance", "", ("指标，万元", "现有数据", "期间与口径")),
    ("metrics", "关键财务指标", ("指标", "数值", "含义")),
    ("finance_conclusion", "财务上的结论", ()),
    ("finance_missing", "申报前必须补的财务数据", ("数据", "当前状态", "用途", "补强时间")),
    ("tasks", "5.1 补强任务表", ("补强动作", "补强价值", "建议节点")),
)
SCORE_TEXT = re.compile(
    r"评分|得分|分数|平台分|质量分|\d+(?:\.\d+)?\s*分(?!钟|贝|米|秒|之)"
    r"|综合评价\s*[:：]\s*\d+(?:\.\d+)?\s*[/／]\s*\d+"
)
PROVINCIAL_PREREQUISITE = re.compile(
    r"前置(?:称号|资格|身份)|科技和创新型中小企业|创新型中小企业|省级科技型中小企业"
)


def screen_research_equipment(fixed_assets_yuan, tiers):
    """Internal recommendation estimate, never an actual equipment-value claim."""
    def amount(value):
        try:
            parsed = Decimal(str(value))
        except InvalidOperation:
            raise ValueError("设备筛选金额必须为非负有限数，单位为元") from None
        if not parsed.is_finite() or parsed < 0:
            raise ValueError("设备筛选金额必须为非负有限数，单位为元")
        return parsed

    checked = []
    for tier in tiers:
        if not tier.get("name") or not tier.get("source"):
            raise ValueError("设备档位须提供名称和已核验政策来源")
        checked.append((amount(tier["equipment_threshold_yuan"]), tier["name"]))
    if fixed_assets_yuan is None:
        return {"status": "missing-fixed-assets", "estimate_yuan": None, "recommended": []}
    estimate = amount(fixed_assets_yuan) * Decimal("0.5")
    matched = [name for threshold, name in sorted(checked) if estimate >= threshold]
    return {"status": "matched" if matched else "no-matching-tier",
            "estimate_yuan": str(estimate), "recommended": matched}


def prepare_sections(fixture):
    sections = deepcopy(validate_sections(fixture))
    screening = fixture.get("research_equipment_screen")
    if screening is None:
        return sections
    result = screen_research_equipment(screening.get("fixed_assets_yuan"), screening["tiers"])
    recommendation = (
        "建议申报方向：" + "、".join(result["recommended"]) + "。完善研发平台建设，支撑专精申报中的持续创新能力。"
        if result["recommended"] else
        "推进企业研发机构建设，统筹研发团队与技术项目，支撑专精申报中的持续创新能力。"
    )
    row = ["企业研发机构", "推荐申报" if result["recommended"] else "优先培育",
           "围绕主导产品建设研发平台", screening["existing_basis"],
           "完善研发组织及运行机制", recommendation]
    rows = sections["next_tasks"].get("rows", [])
    sections["next_tasks"]["rows"] = [r for r in rows if r[0] != "企业研发机构"] + [row]
    return sections


def public_text(value, project_id=None):
    # Filter score clauses, not the facts or qualitative benefits beside them.
    value = value.replace("评分项", "评价指标")
    result = []
    for sentence in re.split(r"(?<=[。；\n])", value):
        clauses = re.findall(r"([^，,]+)([，,]?)", sentence)
        retained = []
        for clause, separator in clauses:
            if SCORE_TEXT.search(clause):
                continue
            if SCORE_TEXT.search(sentence) and re.fullmatch(
                r"\s*(?:获取不到|无法获取|系统自动评定)[。；\n]?", clause
            ):
                continue
            if project_id == "specialized-sme" and PROVINCIAL_PREREQUISITE.search(clause):
                continue
            retained.append(clause + separator)
        text = "".join(retained)
        if len(retained) != len(clauses):
            text = text.rstrip("，,")
            if text and sentence.endswith(("。", "；", "\n")) and not text.endswith(("。", "；", "\n")):
                text += sentence[-1]
        result.append(text)
    return "".join(result).strip()


def public_rows(rows, project_id=None):
    output = []
    for row in rows:
        label = row[0].strip()
        if label == "资格与专注" or SCORE_TEXT.search(label):
            continue
        if label == "发展质量" and any(SCORE_TEXT.search(cell) for cell in row[1:]):
            continue
        if project_id == "specialized-sme" and PROVINCIAL_PREREQUISITE.search(label):
            continue
        cleaned = [public_text(cell, project_id) for cell in row]
        if not all(cleaned):
            raise ValueError(f"{label}去除评分后缺少正文，请补充具体内容")
        output.append(cleaned)
    return output


def validate_sections(fixture):
    if not isinstance(fixture.get("report_date"), str) or not fixture["report_date"].strip():
        raise ValueError("请提供 report_date")
    if fixture.get("project_id") not in {"specialized-sme", "little-giant"}:
        raise ValueError("sme-action 仅适用于专精特新中小企业或小巨人")
    sections = fixture.get("report_sections")
    if not isinstance(sections, dict):
        raise ValueError("请提供 report_sections；不以通用占位词生成成稿")
    for key, _, headers in SECTIONS:
        section = sections.get(key)
        if not isinstance(section, dict):
            raise ValueError(f"report_sections.{key} 缺失")
        paragraphs = section.get("paragraphs")
        if not isinstance(paragraphs, list) or any(
            not isinstance(p, str) or not p.strip() for p in paragraphs
        ):
            raise ValueError(f"{key}.paragraphs 必须包含具体说明")
        rows = section.get("rows", [])
        if not isinstance(rows, list) or any(
            not isinstance(row, list) or len(row) != len(headers)
            or any(not isinstance(cell, str) or not cell.strip() for cell in row)
            for row in rows
        ):
            raise ValueError(f"{key}.rows 必须为 {len(headers)} 列非空文本")
        if key == "tasks" and not rows:
            raise ValueError("请填写 5.1 补强任务表")
        if key in {"product", "peers", "tasks"}:
            visible_rows = public_rows(rows, fixture["project_id"])
            visible_paragraphs = [public_text(p, fixture["project_id"]) for p in paragraphs]
            if not visible_rows and (key == "tasks" or not any(visible_paragraphs)):
                raise ValueError(f"{key} 缺少具体内容，请填写建议、对照或资料缺口说明")
        if not isinstance(section.get("notes", []), list) or any(
            not isinstance(p, str) or not p.strip() for p in section.get("notes", [])
        ):
            raise ValueError(f"{key}.notes 必须为文本数组")
    return sections


def _table(document, headers, rows, project_id=None):
    rows = public_rows(rows, project_id)
    if project_id == "little-giant":
        rows = [([row[0], "已获认定的省级专精特新中小企业", *row[2:]]
                 if len(row) == 2 and row[0] in {"前置资格", "前置称号"}
                 else [row[0], row[1], "已获认定的省级专精特新中小企业", *row[3:]]
                 if len(row) == 6 and row[0] in {"前置资格", "前置称号"}
                 else row) for row in rows]
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    widths = {2: [4, 13], 3: [3.5, 6, 7.5], 4: [3.3, 4.2, 3.1, 6.4],
              6: [2.1, 1.4, 3.3, 2.8, 2.5, 4.9]}[len(headers)]
    if headers[0] == "补强动作":
        widths = [4.5, 10, 2.5]
    if headers[0] == "对标企业或产品体系":
        widths = [3, 4, 5, 5]
    if headers[0] == "主导产品首选":
        widths = [4.25] * 4
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    for index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "EAF0F8" if index == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "D3E0F3")
                borders.append(border)
            cell._tc.get_or_add_tcPr().append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.08
                snap = OxmlElement("w:snapToGrid")
                snap.set(qn("w:val"), "0")
                paragraph._p.get_or_add_pPr().append(snap)
                paragraph.paragraph_format.keep_with_next = index == 0
                for run in paragraph.runs:
                    run.font.size = Pt(8 if len(headers) == 6 else 8.5)
                    run.bold = index == 0
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)


def build_action_document(template_path: Path, fixture):
    sections = prepare_sections(fixture)
    # Retain the distributed master header, footer and brand watermark.
    document = Document(template_path)
    for node in list(document.element.body):
        if node.tag != qn("w:sectPr"):
            document.element.body.remove(node)
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.size = Pt(9)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    normal.element.get_or_add_pPr().append(snap)
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 11)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("000000" if name == "Title" else "233D66")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = False
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        border = style.element.get_or_add_pPr().find(qn("w:pBdr"))
        if border is not None:
            border.getparent().remove(border)
    document.styles["Subtitle"].font.size = Pt(12)
    document.styles["Subtitle"].font.italic = False
    document.styles["Subtitle"].paragraph_format.space_after = Pt(8)
    label = "专精特新中小企业" if fixture["project_id"] == "specialized-sme" else "专精特新小巨人"
    document.add_paragraph(fixture["enterprise"], "Title")
    document.add_paragraph(f"{label}前期评估报告", "Subtitle")
    document.add_paragraph(f"报告日期：{fixture['report_date']}　建议申报年度：{fixture['suggested_year']}\n报告人：共创研究院")
    blocks = {key: (heading, headers) for key, heading, headers in SECTIONS}
    for group_heading, keys in GROUPS:
        group_title = document.add_heading(group_heading, 1)
        group_title.paragraph_format.page_break_before = False
        for key in keys:
            heading, headers = blocks[key]
            if heading:
                section_heading = document.add_heading(heading, 2)
            for paragraph in sections[key]["paragraphs"]:
                if public_text(paragraph, fixture["project_id"]):
                    document.add_paragraph(public_text(paragraph, fixture["project_id"]))
            if sections[key].get("rows"):
                _table(document, headers, sections[key]["rows"], fixture["project_id"])
            for note in sections[key].get("notes", []):
                if public_text(note, fixture["project_id"]):
                    document.add_paragraph(public_text(note, fixture["project_id"]))
    return document
